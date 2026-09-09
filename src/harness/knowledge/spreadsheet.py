"""Spreadsheet parsing for knowledge documents.

WeKnora flattens spreadsheets into ``A: value, B: value`` text, which cannot be
rendered as a table. The original file is downloadable, so the console fetches
it and parses a cell grid here instead.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field

SPREADSHEET_SUFFIXES = (".xlsx", ".xlsm", ".xls", ".csv")


@dataclass(frozen=True)
class SpreadsheetGrid:
    """A bounded, display-oriented view of one spreadsheet sheet."""

    sheet: str
    rows: tuple[tuple[str, ...], ...] = ()
    truncated: bool = False
    extra_sheets: tuple[str, ...] = field(default=())


def is_spreadsheet(filename: str) -> bool:
    return filename.lower().endswith(SPREADSHEET_SUFFIXES)


def _clean(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_spreadsheet(
    filename: str,
    content: bytes,
    *,
    max_rows: int = 200,
    max_columns: int = 40,
) -> SpreadsheetGrid | None:
    """Parse the first sheet of a spreadsheet into a bounded cell grid."""
    lowered = filename.lower()
    if lowered.endswith((".xlsx", ".xlsm")):
        return _parse_xlsx(content, max_rows=max_rows, max_columns=max_columns)
    if lowered.endswith(".xls"):
        return _parse_xls(content, max_rows=max_rows, max_columns=max_columns)
    if lowered.endswith(".csv"):
        return _parse_csv(content, max_rows=max_rows, max_columns=max_columns)
    return None


def _trim(
    rows: list[list[str]], max_rows: int, max_columns: int
) -> tuple[tuple[tuple[str, ...], ...], bool]:
    truncated = False
    trimmed: list[tuple[str, ...]] = []
    for index, row in enumerate(rows):
        if index >= max_rows:
            truncated = True
            break
        if len(row) > max_columns:
            truncated = True
            row = row[:max_columns]
        trimmed.append(tuple(row))
    while trimmed and not any(cell for cell in trimmed[-1]):
        trimmed.pop()
    return tuple(trimmed), truncated


def _parse_xlsx(content: bytes, *, max_rows: int, max_columns: int) -> SpreadsheetGrid | None:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return None
    try:
        names = workbook.sheetnames
        sheet = workbook[names[0]] if names else None
        if sheet is None:
            return None
        rows = [[_clean(cell) for cell in row] for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
    trimmed, truncated = _trim(rows, max_rows, max_columns)
    return SpreadsheetGrid(
        sheet=names[0],
        rows=trimmed,
        truncated=truncated,
        extra_sheets=tuple(names[1:]),
    )


def _parse_xls(content: bytes, *, max_rows: int, max_columns: int) -> SpreadsheetGrid | None:
    import xlrd

    try:
        workbook = xlrd.open_workbook(file_contents=content)
    except Exception:
        return None
    if workbook.nsheets == 0:
        return None
    sheet = workbook.sheet_by_index(0)
    rows = [
        [_clean(sheet.cell_value(row, col)) for col in range(sheet.ncols)]
        for row in range(sheet.nrows)
    ]
    trimmed, truncated = _trim(rows, max_rows, max_columns)
    return SpreadsheetGrid(
        sheet=sheet.name,
        rows=trimmed,
        truncated=truncated,
        extra_sheets=tuple(workbook.sheet_names()[1:]),
    )


def _parse_csv(content: bytes, *, max_rows: int, max_columns: int) -> SpreadsheetGrid | None:
    import csv

    text = content.decode("utf-8", "replace")
    rows = [list(row) for row in csv.reader(io.StringIO(text))]
    trimmed, truncated = _trim(rows, max_rows, max_columns)
    return SpreadsheetGrid(sheet="CSV", rows=trimmed, truncated=truncated)
